"""Thông số móng bản có dầm (sườn): giá trị mặc định, đọc Excel (sheet TONG HOP), kiểm tra, khối lượng.

Khối lượng bê tông, lót, đào, đắp, đệm cát, ván khuôn, vải địa theo công thức sheet "KL be tong" của Excel mẫu.
Trục X theo cạnh L (dài), trục Y theo cạnh B (rộng). 2 dầm mỗi phương, chạy qua tim cổ móng."""
from __future__ import annotations

import math

# ô Excel -> thông số (sheet TONG HOP)
CELLS_NUM = {
    "L1": "K7", "L2": "K8", "B": "K10", "L": "K11",
    "A": "K13",     # rộng dầm
    "HD": "K14",    # cao dầm (tính từ đáy bản)
    "HS": "K16",    # dày bản
    "C": "K17",     # cạnh lỗ mở ở tim móng
    "E": "K18",     # cạnh vát nách chỗ giao dầm
    "BB": "K20",    # cạnh cổ móng
    "H1": "K28", "H2": "K30",
    "HDD": "K36", "HDT": "K37",
    "HC": "K39", "SLOPE": "K40", "MPIT": "K41",
}
CELLS_STR = {"GD": "C5", "COT": "C3", "TENMONG": "C4", "BTM": "K3", "BTL": "K4", "CODE": "K5"}
CELL_HT = "B3"
CELLS_BLN = ("K42", "L42")
# cốt thép: (khoá, ô, hệ số đổi ra mm)
CELLS_REBAR = [
    ("RS_D", "C44", 1), ("RS_S2", "D44", 10), ("RS_S1", "D45", 10),       # bản móng: Ø, @ ô 2 cạnh / ô 4 cạnh (cm)
    ("RCH_D", "I44", 1), ("RCH_S2", "J44", 10), ("RCH_S1", "J45", 10),    # thép chống (cm)
    ("RC_D", "C48", 1), ("RC_N", "D48", 1),                               # thép chủ cổ móng: Ø, số thanh
    ("RT_D", "C49", 1), ("RT_S", "D49", 1), ("RP_D", "C50", 1), ("RP_S", "D50", 1),
    ("RL_D", "C51", 1), ("RL_S", "D51", 1),
    ("RD_D", "I48", 1), ("RD_N", "J48", 1),                               # thép chủ dầm: Ø, số thanh 1 mặt
    ("RDT_D", "I49", 1), ("RDT_S", "J49", 1), ("RDP_D", "I50", 1), ("RDP_S", "J50", 1),
    ("TL_D", "I51", 1), ("TL_S", "J51", 1),                               # thang leo
]
CELLS_COVER = {"COVER_B": "K34", "COVER": "K35"}   # mm trong Excel -> m
CELL_RLAYERS = "F51"                               # "02 lớp"

HT_NAMES = ["ĐỠ", "NÉO GÓC", "NÉO DỪNG"]
SCALES = [50, 75, 100, 125, 150, 200, 250]
SLOPE_DRAW = 0.5          # hệ số mở mái THỂ HIỆN trên bản vẽ (K40 chỉ dùng tính khối lượng)
TL = 0.1                  # dày bê tông lót (m)
ML = 0.1                  # lót tràn ra mỗi bên / vào trong lỗ (m)
NPIER = 4


def default_params() -> dict:
    return {
        "L1": 4.0, "L2": 4.0, "B": 10.0, "L": 12.0,
        "A": 0.5, "HD": 1.0, "HS": 0.35, "C": 1.0, "E": 0.35, "BB": 0.7,
        "H1": 1.0, "H2": 2.5, "HDD": 0.0, "HDT": 0.0,
        "HC": 1.0, "SLOPE": 1.0, "MPIT": 0.6,
        "HT": 0, "BLN": "BLN-M56", "GD": "TKBVTC", "TENMONG": "",
        "SCALE": 100,
        "BTM": "B20", "BTL": "B7.5",
        "COT": "D122-42", "CODE": "",
        "RS_D": 14.0, "RS_S2": 120.0, "RS_S1": 200.0,
        "RCH_D": 10.0, "RCH_S2": 360.0, "RCH_S1": 600.0,
        "RC_D": 25.0, "RC_N": 12.0, "RT_D": 8.0, "RT_S": 200.0, "RP_D": 8.0, "RP_S": 400.0,
        "RL_D": 8.0, "RL_S": 100.0, "RL_N": 2.0,
        "RD_D": 25.0, "RD_N": 8.0, "RDT_D": 8.0, "RDT_S": 200.0, "RDP_D": 8.0, "RDP_S": 400.0,
        "TL_D": 25.0, "TL_S": 300.0,
        "DT_MODE": 0.0,     # đối trọng: 0 = 1/2 diện tích móng, 1 = chọn ô bản
        "DT_SEL": "123456",  # các ô bản đặt đối trọng (khi DT_MODE = 1)
        "RD_ARR": 0.0,      # thép chủ dầm nhiều hàng (hàng dưới > 2 thanh): 0 = các hàng đều nhau, 1 = xếp đầy hàng ngoài
        "COVER": 0.075, "COVER_B": 0.075,
        # khung tên
        "MABV": "",
    }


def type_name(p: dict) -> str:
    if p.get("TENMONG"):
        return p["TENMONG"]
    f = lambda v: f"{v:g}"
    return f"MB{f(p['L1'])}-{f(p['B'])}x{f(p['L'])}-{f(p['H2'])}"


def ht_from_string(s: str) -> int:
    u = s.upper()
    if "DỪNG" in u or "DUNG" in u:
        return 2
    if "NÉO" in u or "NEO" in u or "GÓC" in u or "GOC" in u:
        return 1
    return 0


def _num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def read_excel(path: str, sheet: str = "TONG HOP", p: dict | None = None) -> tuple[dict, list[str]]:
    """Đọc sheet TONG HOP. Trả về (thông số, danh sách cảnh báo)."""
    import openpyxl

    p = dict(p or default_params())
    warn: list[str] = []
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if sheet not in wb.sheetnames:
        raise ValueError(f'Không tìm thấy sheet "{sheet}" trong file Excel.')
    ws = wb[sheet]
    val = lambda a: ws[a].value

    nread = 0
    for k, a in CELLS_NUM.items():
        v = val(a)
        if _num(v):
            p[k] = float(v)
            nread += 1
        elif k in ("HC", "HDD", "HDT") and v in (None, ""):
            p[k] = 0.0
        else:
            warn.append(f"Ô {a} ({k}) không có số, giữ giá trị {p[k]}.")
    for k, a in CELLS_STR.items():
        v = val(a)
        if v not in (None, ""):
            p[k] = str(v).strip()
    v = val(CELL_HT)
    if isinstance(v, str):
        p["HT"] = ht_from_string(v)
    bln = "".join(str(val(a) or "").strip() for a in CELLS_BLN)
    if bln:
        p["BLN"] = bln
    for k, a, f in CELLS_REBAR:
        v = val(a)
        if _num(v) and v > 0:
            p[k] = float(v) * f
        else:
            warn.append(f"Ô {a} ({k}) không có số, giữ giá trị {p[k]:g}.")
    for k, a in CELLS_COVER.items():
        v = val(a)
        if _num(v) and v > 0:
            p[k] = float(v) / 1000.0
    v = val(CELL_RLAYERS)
    if v not in (None, ""):
        try:
            n = int(v) if _num(v) else int("".join(c for c in str(v) if c.isdigit()) or 0)
            if n > 0:
                p["RL_N"] = float(n)
        except ValueError:
            pass
    wb.close()
    if nread == 0:
        raise ValueError("Không đọc được số liệu nào từ sheet. Kiểm tra lại địa chỉ ô.")
    return p, warn


def H(p):
    return p["H1"] + p["H2"]


# ---------------------------------------------------------------------------
# khối lượng (sheet "KL be tong")
def vol_mong(p) -> float:
    b, l, hs, bb, a, hd, c, e = p["B"], p["L"], p["HS"], p["BB"], p["A"], p["HD"], p["C"], p["E"]
    return (b * l * hs + 4 * bb * bb * (H(p) - hd) + 2 * a * (hd - hs) * b + 2 * a * (hd - hs) * (l - 2 * a)
            - c * c * hs + 8 * e * e * (hd - hs))


def vol_above(p) -> float:
    """Thể tích cổ móng trên mặt đất."""
    return 4 * p["BB"] ** 2 * p["H2"]


def vol_lot(p) -> float:
    return (p["B"] + 2 * ML) * (p["L"] + 2 * ML) * TL - max(0.0, p["C"] - 2 * ML) ** 2 * TL


def pit_depth(p) -> float:
    return p["H1"] + TL + p["HC"]


def pit_dims(p):
    """(rộng đáy, dài đáy, rộng miệng, dài miệng) hố đào."""
    u, v = p["B"] + 2 * p["MPIT"], p["L"] + 2 * p["MPIT"]
    d = 2 * pit_depth(p) * p["SLOPE"]
    return u, v, u + d, v + d


def _frustum(h, u, v, w, x):
    return h / 6 * (u * v + w * x + (u + w) * (v + x))


def vol_dao(p) -> float:
    u, v, w, x = pit_dims(p)
    return round(_frustum(pit_depth(p), u, v, w, x), 2)


def vol_cat(p) -> float:
    hc = p["HC"]
    if hc <= 1e-6:
        return 0.0
    u, v = p["B"] + 2 * p["MPIT"], p["L"] + 2 * p["MPIT"]
    d = 2 * hc * p["SLOPE"]
    return round(_frustum(hc, u, v, u + d, v + d), 2)


def vol_dap(p) -> float:
    """Đắp đất hoàn trả hố móng (trừ phần đối trọng nằm dưới mặt đất)."""
    return vol_dao(p) - (vol_mong(p) - vol_above(p)) - vol_lot(p) - vol_cat(p) - vol_dt(p, below=True)


def vol_dd(p) -> float:
    """Đất đắp thêm (đắp móng hđđ): mặt trên B x L, mái 1:1 ra ngoài, trừ cổ móng nằm trong lớp đắp."""
    h = p.get("HDD", 0.0)
    if h <= 1e-6:
        return 0.0
    return _frustum(h, p["B"], p["L"], p["B"] + 2 * h, p["L"] + 2 * h) - NPIER * p["BB"] ** 2 * h


def vol_thua(p) -> float:
    """Đất thừa = đào − đắp hoàn trả − đắp thêm (âm -> 0: phải mua thêm đất)."""
    return max(0.0, vol_dao(p) - vol_dap(p) - vol_dd(p))


# ---------------------------------------------------------------------------
# đối trọng bê tông (hđt): đặt trên mặt bản, ở nửa móng ngược hướng mũi tên phân giác (néo góc) / hướng tuyến
# (néo dừng, theo phương cạnh dài); cột đỡ không có. Ô bản 1..6 (theo nửa móng, u: ra xa tim, v: ngang):
#   1 = trong-trên, 2 = ngoài-trên, 3 = trong-giữa (có lỗ), 4 = ngoài-giữa, 5 = trong-dưới, 6 = ngoài-dưới.
DT_CELLS = {1: ("in", "top"), 2: ("out", "top"), 3: ("in", "mid"), 4: ("out", "mid"), 5: ("in", "bot"),
            6: ("out", "bot")}
DT_MODES = ["1/2 diện tích móng", "Chọn ô bản"]


def dt_orient(p):
    """'x': đối trọng ở nửa +X ; 'y': nửa +Y ; None: không có (cột đỡ hoặc hđt = 0)."""
    ht = int(p.get("HT", 0))
    if p.get("HDT", 0.0) <= 1e-6 or ht == 0:
        return None
    if ht == 2 and p["B"] > p["L"] + 1e-9:
        return "y"
    return "x"


def dt_frame(p, orient=None):
    """(Lu, du, Lv, dv): nửa cạnh và tim dầm theo phương u (ra xa tim móng) và v."""
    orient = orient or dt_orient(p) or "x"
    if orient == "y":
        return p["B"] / 2, p["L2"] / 2, p["L"] / 2, p["L1"] / 2
    return p["L"] / 2, p["L1"] / 2, p["B"] / 2, p["L2"] / 2


def dt_cell_rect(p, num, orient=None):
    """Ô bản num trong hệ (u, v), mép là tim dầm / tim móng / mép móng."""
    Lu, du, Lv, dv = dt_frame(p, orient)
    iu, iv = DT_CELLS[num]
    u0, u1 = (0.0, du) if iu == "in" else (du, Lu)
    v0, v1 = {"top": (dv, Lv), "mid": (-dv, dv), "bot": (-Lv, -dv)}[iv]
    return u0, u1, v0, v1


def dt_parts(p, num):
    """Diện tích các phần của ô bản: S (chỉ có bản), R (dầm + vát nách, ngoài cổ móng), P (cổ móng), H (lỗ)."""
    a, e, bc, c = p["A"], p["E"], p["BB"], p["C"]
    u0, u1, v0, v1 = dt_cell_rect(p, num)
    iu, iv = DT_CELLS[num]
    nu = 2 if iv == "mid" else 1             # số nửa dầm theo phương u nằm trong ô
    nc = 2 if iv == "mid" else 1             # số góc giao dầm (cổ móng, vát) chạm ô
    area = (u1 - u0) * (v1 - v0)
    rib = a / 2 * (v1 - v0) + nu * a / 2 * (u1 - u0) - nc * (a / 2) ** 2
    hau = nc * e * e / 2
    t = bc / 2 - a / 2
    if t <= 0:
        hp = 0.0
    elif 2 * t <= e:
        hp = t * t
    elif t >= e:
        hp = e * e / 2
    else:
        hp = e * e / 2 - (e - t) ** 2
    rp = 2 * (a / 2) * (bc / 2) - (a / 2) ** 2           # phần dầm nằm dưới 1/4 cổ móng
    R = rib + hau - nc * (rp + hp)
    P = nc * (bc / 2) ** 2
    H = c * c / 2 if num == 3 else 0.0
    return dict(A=area, S=area - R - P - H, R=R, P=P, H=H)


def dt_selected(p):
    if dt_orient(p) is None:
        return []
    if int(p.get("DT_MODE", 0)) == 0:
        return list(DT_CELLS)
    return sorted({int(ch) for ch in str(p.get("DT_SEL", "")) if ch in "123456"})


def vol_dt(p, below=False) -> float:
    """Thể tích bê tông đối trọng: phần bản cao hđt, phần trên dầm / vát cao (hđt − (h − h1)), chừa lỗ và cổ móng.
    below = True: chỉ phần nằm dưới mặt đất tự nhiên."""
    h = p.get("HDT", 0.0)
    if below:
        h = min(h, max(0.0, p["H1"] - p["HS"]))
    hr = p["HD"] - p["HS"]
    v = 0.0
    for num in dt_selected(p):
        q = dt_parts(p, num)
        v += q["S"] * h + q["R"] * max(0.0, h - hr)
    return v


def area_vk(p) -> float:
    b, l, hs, bb, a, hd, c, e = p["B"], p["L"], p["HS"], p["BB"], p["A"], p["HD"], p["C"], p["E"]
    hr = hd - hs
    return (2 * (b + l) * hs + 16 * bb * (H(p) - hd) + 4 * c * hs + 16 * math.sqrt(2) * e * hr
            + 4 * (b - 4 * e - 2 * a) * hr + 4 * (l - 4 * e - 2 * a) * hr + 8 * a * hr)


def area_vai(p) -> float:
    """Trải phẳng thành 1 tấm chữ nhật (theo Excel): đáy + 2 mái + 2 dải gấp lên mặt cát tới mép móng + 2 x 0.5."""
    hc = p["HC"]
    if hc <= 1e-6:
        return 0.0
    s = p["SLOPE"]
    u, v = p["B"] + 2 * p["MPIT"], p["L"] + 2 * p["MPIT"]
    sl = math.sqrt(hc * hc + (hc * s) ** 2)
    y5, z5 = u + 2 * hc * s, v + 2 * hc * s
    return (u + 2 * sl + (y5 - p["B"]) + 1) * (v + 2 * sl + (z5 - p["L"]) + 1)


def vol_vua(p) -> float:
    """Vữa M10 chân đế cột: 4 x b² x 0.05."""
    return 4 * p["BB"] ** 2 * 0.05


def bolt_text(p, npier=NPIER) -> str:
    s = p["BLN"]
    for i, c in enumerate(s):
        if c == "M" and i + 1 < len(s) and s[i + 1].isdigit():
            d = ""
            for ch in s[i + 1:]:
                if not ch.isdigit():
                    break
                d += ch
            if d == "48":
                return "8 cặp M48"
            return f"{4 * npier}M{d}"
    return s


# ---------------------------------------------------------------------------
def validate(p) -> str | None:
    keys = ("L1", "L2", "B", "L", "A", "HD", "HS", "BB", "H1", "H2")
    if min(p[k] for k in keys) < 1e-6:
        return "Có kích thước <= 0. Kiểm tra lại các ô nhập."
    if p["C"] < 0 or p["E"] < 0:
        return "Lỗ mở c và vát nách e không được âm (nhập 0 nếu không có)."
    if p["HC"] < 0:
        return "Dày đệm cát không được âm (nhập 0 nếu không có)."
    if p["HDD"] < 0 or p["HDT"] < 0:
        return "Chiều cao đắp móng hđđ và đối trọng hđt không được âm (nhập 0 nếu không có)."
    if p["HDD"] >= p["H2"] - 1e-6:
        return "Chiều cao đắp móng hđđ phải nhỏ hơn H2 (cổ móng phải nhô lên khỏi mặt đắp)."
    if p["SLOPE"] < 1e-6:
        return "Hệ số mở mái phải > 0."
    if p["MPIT"] < ML:
        return "Mở rộng đáy hố đào phải >= lót ra mỗi bên (0.1 m)."
    if p["HD"] <= p["HS"]:
        return "Cao dầm h phải lớn hơn dày bản h1."
    if H(p) <= p["HD"]:
        return "H = H1 + H2 phải lớn hơn cao dầm h."
    if p["BB"] < p["A"]:
        return "Cạnh cổ móng b phải >= rộng dầm a."
    for ln, l1, nm in ((p["L"], p["L1"], "L"), (p["B"], p["L2"], "B")):
        if ln < l1 + p["A"] + 2 * p["E"] + 0.2:
            return f"{nm} quá nhỏ so với khoảng cách tim dầm và vát nách."
        if l1 - p["A"] - 2 * p["E"] < p["C"] + 0.2:
            return "Lỗ mở c không lọt giữa 2 dầm (trừ vát nách)."
    keys = ("RS_D", "RS_S1", "RS_S2", "RCH_D", "RCH_S1", "RCH_S2", "RC_D", "RT_D", "RT_S", "RP_D", "RP_S",
            "RL_D", "RL_S", "RD_D", "RD_N", "RDT_D", "RDT_S", "RDP_D", "RDP_S", "TL_D", "TL_S")
    if min(p[k] for k in keys) < 1e-6:
        return "Đường kính / khoảng cách / số thanh thép phải > 0."
    n = int(round(p["RC_N"]))
    if n < 4 or n % 4:
        return "Số thanh thép chủ cổ móng phải là bội số của 4 (>= 4)."
    if p["RL_N"] < 1:
        return "Số lớp lưới thép đầu trụ phải >= 1."
    if p["COVER"] < 0 or p["COVER_B"] < 0:
        return "Lớp BT bảo vệ không được âm."
    return None


GAP_MIN = 0.04      # khoảng hở ngang tối thiểu giữa các thanh thép chủ (m)


ARR_NAMES = ["Hàng đều nhau", "Đầy hàng ngoài"]


def rib_rows(n, m, arr):
    """Số thanh mỗi hàng thép chủ dầm (1 mặt), m = số thanh tối đa 1 hàng (khoảng hở ngang >= 40 mm).
    Xếp đầy hàng ngoài trước, hàng dưới tối thiểu 2 thanh (2 bên); nếu hàng dưới > 2 thanh và chọn
    'đều nhau' thì chia đều các hàng."""
    n, m = max(1, int(n)), max(2, int(m))
    nrow = -(-n // m)
    fill = [m] * (nrow - 1) + [n - m * (nrow - 1)]
    if nrow > 1 and fill[-1] < 2 and fill[-2] - (2 - fill[-1]) >= 2:   # hàng dưới tối thiểu 2 thanh ở 2 bên
        fill[-2] -= 2 - fill[-1]
        fill[-1] = 2
    if fill[-1] <= 2 or int(arr) == 1:
        return fill
    per = [-(-n // nrow)] * nrow
    per[-1] = n - sum(per[:-1])
    return per


def warnings(p) -> list:
    """Cảnh báo (không chặn dựng bản vẽ): khoảng hở ngang giữa các thanh thép chủ < 40 mm."""
    out = []
    if p.get("HDT", 0.0) > 1e-6:
        if int(p.get("HT", 0)) == 0:
            out.append("Cột đỡ không bố trí đối trọng: bỏ qua hđt.")
        elif not dt_selected(p):
            out.append("Chưa chọn ô bản nào để đặt đối trọng.")
        elif p["HS"] + p["HDT"] > p["H1"] + 1e-6:
            out.append("Đối trọng nhô lên khỏi mặt đất tự nhiên.")
    try:
        m = 1 + int(round(p["RC_N"])) // 4
        d = p["RC_D"] / 1000.0
        xi = p["BB"] / 2 - p["COVER"] - d / 2
        if m > 1:
            gap = 2 * xi / (m - 1) - d
            if gap < GAP_MIN - 1e-9:
                out.append(f"Khoảng hở thép chủ cổ móng {gap * 1000:.0f} mm < {GAP_MIN * 1000:.0f} mm.")
        d = p["RD_D"] / 1000.0
        w_in = p["A"] - 2 * p["COVER"] - 2 * p["RDT_D"] / 1000.0
        if p["RD_N"] >= 2 and w_in - 2 * d < GAP_MIN - 1e-9:
            out.append(f"Dầm rộng {p['A'] * 1000:.0f} mm không đặt được 2 thanh Ø{p['RD_D']:.0f} "
                       f"với khoảng hở >= {GAP_MIN * 1000:.0f} mm.")
    except (KeyError, TypeError, ValueError, ZeroDivisionError):
        pass
    return out


def auto_scale(p) -> int:
    """Tỷ lệ mặt cắt A-A và mặt bằng (tờ 1)."""
    wx = p["L"] + 2 * p["MPIT"] + 2 * SLOPE_DRAW * pit_depth(p)
    hy = H(p) + TL + p["HC"] + p["B"] + 2 * ML
    k = min(215.0 / wx, 190.0 / hy)
    for d in SCALES:
        if 1000.0 / d <= k:
            return d
    return SCALES[-1]


def is_detail(p) -> bool:
    return "BVTC" in str(p.get("GD", "")).upper()
