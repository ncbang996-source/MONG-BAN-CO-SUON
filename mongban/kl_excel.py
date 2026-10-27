"""Xuất bảng khối lượng ra Excel, mọi khối lượng tính bằng công thức tham chiếu sheet THONG SO."""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName

from . import model
from .geom import KGM, Geo, tie_plan_mg

THIN = Side(style="thin", color="808080")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill("solid", fgColor="DDEBF7")
INPUT = PatternFill("solid", fgColor="FFF2CC")
B = Font(bold=True)

# (tên, ký hiệu, khoá thông số, hệ số đổi, đơn vị)
# Lưu ý: tên định danh Excel không phân biệt hoa/thường -> cổ móng là TS_bc (tránh trùng TS_B)
PARAMS = [
    ("Tim chân cột theo X", "L1", "L1", 1, "m"),
    ("Tim chân cột theo Y", "L2", "L2", 1, "m"),
    ("Chiều dài móng (theo X)", "L", "L", 1, "m"),
    ("Chiều rộng móng (theo Y)", "B", "B", 1, "m"),
    ("Dày bản móng", "h1", "HS", 1, "m"),
    ("Rộng dầm", "a", "A", 1, "m"),
    ("Cao dầm (tính từ đáy bản)", "h", "HD", 1, "m"),
    ("Cạnh vát nách", "e", "E", 1, "m"),
    ("Cạnh lỗ mở tim móng", "c", "C", 1, "m"),
    ("Cạnh cổ móng", "b", "BB", 1, "m"),
    ("Chiều sâu chôn móng", "H1", "H1", 1, "m"),
    ("Chiều cao trên mặt đất", "H2", "H2", 1, "m"),
    ("Dày bê tông lót", "tl", None, model.TL, "m"),
    ("Lót tràn ra mỗi bên", "ml", None, model.ML, "m"),
    ("Dày đệm cát", "hc", "HC", 1, "m"),
    ("Hệ số mở mái hố đào", "m", "SLOPE", 1, ""),
    ("Mở rộng đáy hố đào (mỗi bên)", "mp", "MPIT", 1, "m"),
    ("Số thân trụ", "n", None, 4, "trụ"),
]
# thông số cốt thép (đơn vị như bảng nhập: Ø, @ theo mm)
PARAMS_S = [
    ("Lớp BT bảo vệ đáy móng", "cvb", "COVER_B", 1, "m"),
    ("Lớp BT bảo vệ còn lại", "cv", "COVER", 1, "m"),
    ("Thép bản: Ø", "ds", "RS_D", 1, "mm"),
    ("Thép bản: @ ô 4 cạnh ngàm (ô giữa)", "s1", "RS_S1", 1, "mm"),
    ("Thép bản: @ ô 2 cạnh ngàm (ô biên)", "s2", "RS_S2", 1, "mm"),
    ("Thép chống: Ø", "dch", "RCH_D", 1, "mm"),
    ("Thép chống: @ ô 4 cạnh ngàm", "sch1", "RCH_S1", 1, "mm"),
    ("Thép chống: @ ô 2 cạnh ngàm", "sch2", "RCH_S2", 1, "mm"),
    ("Thép chủ dầm: Ø", "dm", "RD_D", 1, "mm"),
    ("Thép chủ dầm: số thanh 1 mặt", "nd", "RD_N", 1, "thanh"),
    ("Đai dầm: Ø", "dt", "RDT_D", 1, "mm"),
    ("Đai dầm: @", "st", "RDT_S", 1, "mm"),
    ("Đai phụ dầm: Ø", "dp", "RDP_D", 1, "mm"),
    ("Đai phụ dầm: @", "sp", "RDP_S", 1, "mm"),
    ("Thép chủ cổ móng: Ø", "dcm", "RC_D", 1, "mm"),
    ("Thép chủ cổ móng: số thanh 1 cổ", "ncm", "RC_N", 1, "thanh"),
    ("Đai cổ móng: Ø", "dh", "RT_D", 1, "mm"),
    ("Đai cổ móng: @", "sh", "RT_S", 1, "mm"),
    ("Đai phụ cổ móng: Ø", "dph", "RP_D", 1, "mm"),
    ("Đai phụ cổ móng: @", "sph", "RP_S", 1, "mm"),
    ("Thép chủ dầm nhiều hàng: 0 = các hàng đều nhau, 1 = đầy hàng ngoài", "arr", "RD_ARR", 1, ""),
    ("Lưới đầu trụ: Ø", "dl", "RL_D", 1, "mm"),
    ("Lưới đầu trụ: @", "sl", "RL_S", 1, "mm"),
    ("Lưới đầu trụ: số lớp", "nl", "RL_N", 1, "lớp"),
    ("Thang leo: Ø", "dtl", "TL_D", 1, "mm"),
    ("Thang leo: khoảng cách bậc", "stl", "TL_S", 1, "mm"),
    ("Móc đai, đai phụ", "hk", None, 0.06, "m"),
    ("Chân thanh chống", "lch", None, 0.1, "m"),
    ("Đoạn neo chân thép chủ cổ móng", "lneo", None, 0.5, "m"),
    ("Thang leo: chân ngàm vào cổ", "ltl", None, 0.35, "m"),
    ("Thang leo: bề rộng bậc", "wtl", None, 0.3, "m"),
    ("Thang leo bắt đầu từ cao độ (so với mặt đất)", "htl", None, 1.5, "m"),
]
NAME = {"L1": "TS_L1", "L2": "TS_L2", "L": "TS_L", "B": "TS_B", "h1": "TS_hs", "a": "TS_a", "h": "TS_hd",
        "e": "TS_e", "c": "TS_c", "b": "TS_bc", "H1": "TS_H1", "H2": "TS_H2", "tl": "TS_tl", "ml": "TS_ml",
        "hc": "TS_hc", "m": "TS_m", "mp": "TS_mp", "n": "TS_n"}


NAMES = {1: "Thép bản móng theo Y", 2: "Thép bản theo Y, ngắt tại lỗ", 3: "Thép bản móng theo X",
         4: "Thép bản theo X, ngắt tại lỗ", 5: "Thép xiên góc lỗ", 6: "Thanh chống", 7: "Thép vát nách (dài)",
         8: "Thép vát nách (ngắn)", 9: "Đai vuông vát nách", 10: "Thép chủ dầm theo Y (dài suốt)",
         11: "Thép chủ dầm theo Y (ngắt)", 12: "Thép giá dầm theo Y", 13: "Thép chủ dầm theo X (dài suốt)",
         14: "Thép chủ dầm theo X (ngắt)", 15: "Thép giá dầm theo X", 16: "Đai dầm", 17: "Đai phụ ngang dầm",
         18: "Đai phụ đứng dầm", 19: "Thép chủ cổ móng", 20: "Đai cổ móng", 21: "Đai phụ thẳng cổ móng",
         22: "Đai bát giác cổ móng", 23: "Lưới thép đầu trụ", 24: "Thang leo"}


def _cell(ws, r, c, v, fmt=None, bold=False, fill=None, align=None):
    x = ws.cell(row=r, column=c, value=v)
    x.border = BOX
    if fmt:
        x.number_format = fmt
    if bold:
        x.font = B
    if fill:
        x.fill = fill
    x.alignment = align or Alignment(vertical="center", wrap_text=True)
    return x


def export(p: dict, g: Geo, path: str):
    if not getattr(g, "kl", None):
        g.calc_kl()
    wb = Workbook()
    ws = wb.active
    ws.title = "THONG SO"
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 8
    ws["A1"] = f"THÔNG SỐ MÓNG {model.type_name(p)} - CỘT {model.HT_NAMES[int(p['HT'])]} {p['COT']} ({p['GD']})"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Ô nền vàng là số liệu đầu vào - sửa ở đây, sheet KHOI LUONG tự tính lại."
    ws["A2"].font = Font(italic=True, color="808080")
    r = 4
    for i, h in enumerate(["STT", "Thông số", "Ký hiệu", "Giá trị", "Đơn vị"], 1):
        _cell(ws, r, i, h, bold=True, fill=HEAD, align=Alignment(horizontal="center", vertical="center"))
    for k, (nm, sym, key, f, unit) in enumerate(PARAMS, 1):
        r += 1
        v = p[key] * f if key else f
        _cell(ws, r, 1, k, align=Alignment(horizontal="center"))
        _cell(ws, r, 2, nm)
        _cell(ws, r, 3, sym, align=Alignment(horizontal="center"))
        _cell(ws, r, 4, v, "0.000", fill=INPUT)
        _cell(ws, r, 5, unit, align=Alignment(horizontal="center"))
        wb.defined_names[NAME[sym]] = DefinedName(NAME[sym], attr_text=f"'THONG SO'!$D${r}")
    r += 1
    c_ = _cell(ws, r + 1, 1, "THÔNG SỐ CỐT THÉP", bold=True, fill=HEAD)
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=5)
    c_.alignment = Alignment(vertical="center")
    r += 1
    for k, (nm, sym, key, f, unit) in enumerate(PARAMS_S, len(PARAMS) + 1):
        r += 1
        v = p.get(key, 0.0) * f if key else f
        _cell(ws, r, 1, k, align=Alignment(horizontal="center"))
        _cell(ws, r, 2, nm)
        _cell(ws, r, 3, sym, align=Alignment(horizontal="center"))
        _cell(ws, r, 4, v, "0.000" if unit == "m" else "0", fill=INPUT)
        _cell(ws, r, 5, unit, align=Alignment(horizontal="center"))
        wb.defined_names[f"TS_{sym}"] = DefinedName(f"TS_{sym}", attr_text=f"'THONG SO'!$D${r}")
    # bảng tra trọng lượng thép
    r += 2
    ws.cell(row=r, column=2, value="TRỌNG LƯỢNG THÉP (TCVN 1651-2018)").font = B
    r += 1
    _cell(ws, r, 2, "Đường kính (mm)", bold=True, fill=HEAD)
    _cell(ws, r, 3, "kg/m", bold=True, fill=HEAD)
    r0 = r + 1
    for d, w in sorted(KGM.items()):
        r += 1
        _cell(ws, r, 2, d)
        _cell(ws, r, 3, w, "0.000", fill=INPUT)
    wb.defined_names["TS_KGM"] = DefinedName("TS_KGM", attr_text=f"'THONG SO'!$B${r0}:$C${r}")
    # bảng tra bố trí đai phụ cổ móng theo số thanh 1 cạnh (m) và bước giữ thanh g (1 hoặc 2)
    r += 2
    ws.cell(row=r, column=2, value="BỐ TRÍ ĐAI PHỤ CỔ MÓNG (chương trình lập, dùng cho sheet KHOI LUONG)").font = B
    r += 1
    for i, h in enumerate(["Số thanh 1 cạnh m", "g=1: số đai thẳng", "g=1: chỉ số max đai bát giác",
                           "g=2: số đai thẳng", "g=2: chỉ số max đai bát giác"], 2):
        _cell(ws, r, i, h, bold=True, fill=HEAD)
    r0 = r + 1
    for m in range(2, 31):
        r += 1
        _cell(ws, r, 2, m)
        for k, gg in enumerate((1, 2)):
            grp, st = tie_plan_mg(m, gg)
            _cell(ws, r, 3 + 2 * k, len(st))
            _cell(ws, r, 4 + 2 * k, max(grp) if grp else -1)
    wb.defined_names["TS_DPC"] = DefinedName("TS_DPC", attr_text=f"'THONG SO'!$B${r0}:$F${r}")

    # ------------------------------------------------------------------ KHOI LUONG
    wk = wb.create_sheet("KHOI LUONG")
    for col, w in zip("ABCDEFGHIJ", (6, 44, 10, 70, 12, 9, 12, 12, 34, 40)):
        wk.column_dimensions[col].width = w
    wk["A1"] = f"BẢNG KHỐI LƯỢNG CHO 01 VỊ TRÍ MÓNG {model.type_name(p)} - CỘT {p['COT']}"
    wk["A1"].font = Font(bold=True, size=12)
    r = 3
    for i, h in enumerate(["STT", "Hạng mục", "Ký hiệu", "Công thức / diễn giải", "Khối lượng", "Đơn vị"], 1):
        _cell(wk, r, i, h, bold=True, fill=HEAD, align=Alignment(horizontal="center", vertical="center"))
    rows = {}

    def line(stt, name, sym, formula, desc, unit, bold=False):
        nonlocal r
        r += 1
        _cell(wk, r, 1, stt, align=Alignment(horizontal="center"))
        _cell(wk, r, 2, name, bold=bold)
        _cell(wk, r, 3, sym, align=Alignment(horizontal="center"))
        _cell(wk, r, 4, desc)
        _cell(wk, r, 5, formula, "0.00", bold=bold)
        _cell(wk, r, 6, unit, align=Alignment(horizontal="center"))
        if sym:
            rows[sym] = f"$E${r}"
        return f"$E${r}"

    def sec(t):
        nonlocal r
        r += 1
        c = _cell(wk, r, 1, t, bold=True, fill=HEAD)
        wk.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        c.alignment = Alignment(vertical="center")

    sec("A. KÍCH THƯỚC TRUNG GIAN")
    hr = line("", "Phần dầm nhô trên bản", "hr", "=TS_hd-TS_hs", "h − h1", "m")
    hd = line("", "Chiều sâu hố đào", "Hd", "=TS_H1+TS_tl+TS_hc", "H1 + dày lót + dày đệm cát", "m")
    a1 = line("", "Đáy hố đào theo X", "a1", "=TS_L+2*TS_mp", "L + 2·mp", "m")
    b1 = line("", "Đáy hố đào theo Y", "b1", "=TS_B+2*TS_mp", "B + 2·mp", "m")
    a2 = line("", "Miệng hố đào theo X", "a2", f"={a1}+2*TS_m*{hd}", "a1 + 2·m·Hd", "m")
    b2 = line("", "Miệng hố đào theo Y", "b2", f"={b1}+2*TS_m*{hd}", "b1 + 2·m·Hd", "m")
    dc = line("", "Mở rộng mái trong lớp đệm cát", "dc", "=TS_m*TS_hc", "m · hc", "m")

    sec("B. BÊ TÔNG, ĐỆM CÁT, VẢI ĐỊA")
    vb = line(1, f"Bê tông móng {p['BTM']}", "Vbt",
              f"=TS_B*TS_L*TS_hs+TS_n*TS_bc^2*(TS_H1+TS_H2-TS_hd)+2*TS_a*{hr}*TS_B+2*TS_a*{hr}*(TS_L-2*TS_a)"
              f"-TS_c^2*TS_hs+8*TS_e^2*{hr}",
              "B·L·h1 + n·b²·(H − h) + 2a(h−h1)·B + 2a(h−h1)·(L − 2a) − c²·h1 + 8e²(h−h1)", "m3", True)
    vl = line(2, f"Bê tông lót {p['BTL']}", "Vlot",
              "=(TS_B+2*TS_ml)*(TS_L+2*TS_ml)*TS_tl-MAX(0,TS_c-2*TS_ml)^2*TS_tl",
              "(B + 2·ml)·(L + 2·ml)·tl − (c − 2·ml)²·tl", "m3", True)
    vc = line(3, "Đệm cát", "Vcat",
              f"=IF(TS_hc>0,ROUND(TS_hc/6*({a1}*{b1}+({a1}+2*{dc})*({b1}+2*{dc})+(2*{a1}+2*{dc})*(2*{b1}+2*{dc})),2),0)",
              "Hình chóp cụt: hc/6·(A1 + A2 + (a1+a2)(b1+b2))", "m3", True)
    line(4, "Vải địa kỹ thuật", "Svai",
         f"=IF(TS_hc>0,({b1}+2*SQRT(TS_hc^2+(TS_hc*TS_m)^2)+({b1}+2*{dc}-TS_B)+1)"
         f"*({a1}+2*SQRT(TS_hc^2+(TS_hc*TS_m)^2)+({a1}+2*{dc}-TS_L)+1),0)",
         "Trải phẳng 1 tấm: (đáy + 2 mái + 2 dải gấp tới mép móng + 1) theo mỗi phương", "m2", True)
    line("", "Vữa M10 chân đế cột", "Vvua", "=TS_n*TS_bc^2*0.05", "n · b² · 0.05", "m3")

    sec("C. ĐẤT, VÁN KHUÔN")
    vd = line(5, "Đào đất hố móng", "Vdao",
              f"=ROUND({hd}/6*({a1}*{b1}+{a2}*{b2}+({a1}+{a2})*({b1}+{b2})),2)",
              "Hình chóp cụt: Hd/6·(a1·b1 + a2·b2 + (a1+a2)(b1+b2))", "m3", True)
    vu = line("", "Thể tích móng nằm dưới mặt đất", "Vmd", f"={vb}-TS_n*TS_bc^2*TS_H2", "Vbt − n·b²·H2", "m3")
    vdp = line(6, "Đắp đất hoàn trả hố móng", "Vdap", f"={vd}-{vu}-{vl}-{vc}", "Vdao − Vmd − Vlot − Vcat", "m3", True)
    line(7, "Đất thừa (đắp ụ quanh cổ móng)", "Vthua", f"={vd}-{vdp}", "Vdao − Vdap", "m3", True)
    line(8, "Ván khuôn móng", "VK",
         f"=2*(TS_B+TS_L)*TS_hs+4*TS_n*TS_bc*(TS_H1+TS_H2-TS_hd)+4*TS_c*TS_hs+16*SQRT(2)*TS_e*{hr}"
         f"+4*(TS_B-4*TS_e-2*TS_a)*{hr}+4*(TS_L-4*TS_e-2*TS_a)*{hr}+8*TS_a*{hr}",
         "Mép bản + cổ móng (trên dầm) + lỗ mở + mặt vát + thành dầm + đầu dầm", "m2", True)

    sec("D. CỐT THÉP - ĐẠI LƯỢNG TRUNG GIAN (công thức theo thông số, giống bảng thống kê trên bản vẽ)")

    def LIN(w, sp):          # rải thép: số thanh = ROUND(w/@) + 1 (>= 2), w = đoạn rải đã trừ lớp bảo vệ
        return f"IF(({w})<0,0,IF(({w})<=0.000001,1,MAX(2,ROUND(({w})/({sp}),0)+1)))"

    def CELLS(w, sp):        # thanh chống: số ô = ROUNDUP(w/@)
        return f"IF(({w})<=0.000001,0,MAX(1,ROUNDUP(({w})/({sp})-0.000002,0)))"

    def STA(w, sp):          # đai rải đều, 2 đầu có đai: 1 + ROUNDUP(w/@)
        return f"IF(({w})<=0.000001,0,1+MAX(1,ROUNDUP(({w})/({sp})-0.000002,0)))"
    has_hole = "TS_c>0.000001"
    d1 = line("", "Mép dầm tới mép móng theo X", "D1", "=TS_L/2-TS_L1/2-TS_a/2", "L/2 − L1/2 − a/2", "m")
    d2 = line("", "Mép dầm tới mép móng theo Y", "D2", "=TS_B/2-TS_L2/2-TS_a/2", "B/2 − L2/2 − a/2", "m")
    wxi = line("", "Ô giữa theo X (mỗi bên lỗ)", "wx", f"=IF({has_hole},TS_L1/2-TS_a/2-TS_c/2,TS_L1-TS_a)",
               "Có lỗ: L1/2 − a/2 − c/2 (2 ô); không lỗ: L1 − a", "m")
    wyi = line("", "Ô giữa theo Y (mỗi bên lỗ)", "wy", f"=IF({has_hole},TS_L2/2-TS_a/2-TS_c/2,TS_L2-TS_a)",
               "Có lỗ: L2/2 − a/2 − c/2 (2 ô); không lỗ: L2 − a", "m")
    k_in = f"IF({has_hole},2,1)"
    nx = line("", "Số thanh thép bản theo Y (1 lớp)", "nX",
              f"=2*{LIN(f'{d1}-2*TS_cv', 'TS_s2/1000')}+{k_in}*{LIN(f'{wxi}-2*TS_cv', 'TS_s1/1000')}",
              "2 ô biên: ROUND((D1 − 2cv)/@2)+1 ; ô giữa: ROUND((wx − 2cv)/@1)+1", "thanh")
    ny = line("", "Số thanh thép bản theo X (1 lớp)", "nY",
              f"=2*{LIN(f'{d2}-2*TS_cv', 'TS_s2/1000')}+{k_in}*{LIN(f'{wyi}-2*TS_cv', 'TS_s1/1000')}",
              "2 ô biên: ROUND((D2 − 2cv)/@2)+1 ; ô giữa: ROUND((wy − 2cv)/@1)+1", "thanh")
    nh = line("", "Số thanh qua phạm vi lỗ (mỗi phương, 1 lớp)", "nh",
              f"=IF({has_hole},{LIN('TS_c-2*TS_cv', 'TS_s1/1000')},0)", "ROUND((c − 2cv)/@1)+1", "thanh")
    chx = line("", "Số hàng thanh chống theo X", "chX",
               f"=2*{CELLS(f'{d1}-2*TS_cv', 'TS_sch2/1000')}+{k_in}*{CELLS(f'{wxi}-2*TS_cv', 'TS_sch1/1000')}"
               f"+IF({has_hole},{CELLS('TS_c-2*TS_cv', 'TS_sch1/1000')},0)",
               "Mỗi ô: ROUNDUP((bề rộng ô − 2cv)/@ chống) ; cộng ô lỗ", "hàng")
    chy = line("", "Số hàng thanh chống theo Y", "chY",
               f"=2*{CELLS(f'{d2}-2*TS_cv', 'TS_sch2/1000')}+{k_in}*{CELLS(f'{wyi}-2*TS_cv', 'TS_sch1/1000')}"
               f"+IF({has_hole},{CELLS('TS_c-2*TS_cv', 'TS_sch1/1000')},0)",
               "Mỗi ô: ROUNDUP((bề rộng ô − 2cv)/@ chống) ; cộng ô lỗ", "hàng")
    chh = line("", "Thanh chống rơi vào lỗ (bỏ)", "chH",
               f"=IF({has_hole},{CELLS('TS_c-2*TS_cv', 'TS_sch1/1000')}^2,0)", "(số ô trong lỗ)²", "thanh")
    nf = line("", "Thép chủ dầm dài suốt (1 mặt)", "nf",
              "=IF(TS_nd<5,TS_nd,MIN(TS_nd,IF(MOD(TS_nd,4)=0,TS_nd/2,INT(TS_nd/2)+1)))",
              "n < 5: không cắt (= n) ; n chia hết 4: n/2 ; còn lại: INT(n/2)+1", "thanh")
    ncut = line("", "Thép chủ dầm ngắt (1 mặt)", "nc", f"=TS_nd-{nf}", "n − nf", "thanh")
    cuty = line("", "Đoạn cắt mỗi đầu thanh ngắt - dầm theo Y", "cutY", f"=ROUNDUP({d2}/0.2-0.000001,0)*0.1",
                "½·D2 làm tròn lên 100 mm (tính từ mép móng)", "m")
    cutx = line("", "Đoạn cắt mỗi đầu thanh ngắt - dầm theo X", "cutX", f"=ROUNDUP({d1}/0.2-0.000001,0)*0.1",
                "½·D1 làm tròn lên 100 mm (tính từ mép móng)", "m")
    hg = line("", "Khoảng hở ngang tối thiểu thép chủ dầm", "hgap", "=MAX(0.04,TS_dm/1000)", "max(40 mm, Ø)", "m")
    mrow = line("", "Số thanh tối đa 1 hàng", "m1",
                f"=MAX(2,INT((TS_a-2*TS_cv-2*TS_dt/1000+{hg}+0.000000001)/(TS_dm/1000+{hg})))",
                "INT((a − 2cv − 2Øđai + hở)/(Ø + hở))", "thanh")
    nrow = line("", "Số hàng thép chủ dầm (1 mặt)", "nrow", f"=ROUNDUP(TS_nd/{mrow},0)", "ROUNDUP(n/m1)", "hàng")
    gv = line("", "Khoảng hở giữa các hàng", "gapv", "=MAX(0.025,TS_dm/1000)", "max(25 mm, Ø)", "m")
    zb = line("", "Cao độ hàng trong cùng - mặt dưới", "z0",
              f"=TS_cvb+MAX(TS_dt/1000,2*TS_ds/1000)+TS_dm/2000+({nrow}-1)*(TS_dm/1000+{gv})",
              "cvb + max(Øđai, 2Ø bản) + Ø/2 + (nrow−1)(Ø + hở)", "m")
    zt = line("", "Cao độ hàng trong cùng - mặt trên", "z1",
              f"=TS_hd-TS_cv-TS_dt/1000-TS_dm/2000-({nrow}-1)*(TS_dm/1000+{gv})",
              "h − cv − Øđai − Ø/2 − (nrow−1)(Ø + hở)", "m")
    ngia = line("", "Số mức thép giá mỗi bên (TCVN 5574:2018 10.3.3.3)", "ngia",
                f"=IF(TS_hd>0.700001,MAX(0,ROUNDUP(({zt}-{zb})/0.4-0.000001,0)-1),0)",
                "h > 700: ROUNDUP((z1 − z0)/400) − 1 ; h <= 700: 0", "mức")
    wq = "(TS_a/2+TS_e)"
    sty = line("", "Số đai dầm 1 dầm theo Y", "nstY",
               f"=2*{STA(f'TS_B/2-TS_cv-TS_L2/2-{wq}', 'TS_st/1000')}+{STA(f'TS_L2-2*{wq}', 'TS_st/1000')}",
               "Bỏ vùng giao dầm + vát: 2 đoạn đầu (B/2 − cv − L2/2 − a/2 − e) + đoạn giữa (L2 − a − 2e); "
               "mỗi đoạn 1 + ROUNDUP(dài/@)", "đai")
    stx = line("", "Số đai dầm 1 dầm theo X", "nstX",
               f"=2*{STA(f'TS_L/2-TS_cv-TS_L1/2-{wq}', 'TS_st/1000')}+{STA(f'TS_L1-2*{wq}', 'TS_st/1000')}",
               "Như trên, theo L, L1", "đai")
    tiy = line("", "Số vị trí đai phụ 1 dầm theo Y", "ntY",
               f"=2*{STA(f'TS_B/2-TS_cv-TS_L2/2-{wq}', 'TS_sp/1000')}+{STA(f'TS_L2-2*{wq}', 'TS_sp/1000')}",
               "Như đai dầm, theo @ đai phụ", "vị trí")
    tix = line("", "Số vị trí đai phụ 1 dầm theo X", "ntX",
               f"=2*{STA(f'TS_L/2-TS_cv-TS_L1/2-{wq}', 'TS_sp/1000')}+{STA(f'TS_L1-2*{wq}', 'TS_sp/1000')}",
               "Như đai dầm, theo @ đai phụ", "vị trí")
    lastf = line("", "Số thanh hàng cuối khi xếp đầy", "kL0", f"=TS_nd-{mrow}*({nrow}-1)", "n − m1·(nrow − 1)", "thanh")
    adj = line("", "Bù để hàng dưới tối thiểu 2 thanh (2 bên)", "adj",
               f"=IF(AND({nrow}>1,{lastf}<2,{mrow}-(2-{lastf})>=2),2-{lastf},0)",
               "Hàng cuối < 2 thanh: chuyển bớt từ hàng trên xuống", "thanh")
    fm = line("", "Xếp đầy hàng ngoài (1) / các hàng đều nhau (0)", "fill",
              f"=IF(OR({lastf}+{adj}<=2,TS_arr=1),1,0)", "Hàng cuối <= 2 thanh luôn xếp đầy; còn lại theo lựa chọn arr", "")
    k1 = line("", "Số thanh hàng ngoài cùng (1 mặt)", "k1",
              f"=IF({nrow}<=1,TS_nd,IF({fm}=1,IF({nrow}=2,{mrow}-{adj},{mrow}),ROUNDUP(TS_nd/{nrow},0)))",
              "Xếp đầy: m1 (trừ phần bù) ; đều nhau: ROUNDUP(n/nrow)", "thanh")
    k2 = line("", "Số thanh hàng 2 (1 mặt)", "k2",
              f"=IF({nrow}<2,0,IF({fm}=1,IF({nrow}=2,{lastf}+{adj},IF({nrow}=3,{mrow}-{adj},{mrow})),"
              f"IF({nrow}=2,TS_nd-ROUNDUP(TS_nd/2,0),ROUNDUP(TS_nd/{nrow},0))))",
              "Theo cách xếp ở trên", "thanh")
    has2 = line("", "Có đai phụ giữ lớp 2", "t2", f"=IF({k2}>2,1,0)",
                "= 1 khi hàng 2 có > 2 thanh (hàng 2 chỉ 2 thanh thì nằm sát 2 nhánh đai chính)", "")
    full2 = line("", "Đai phụ giữ lớp 2 rải suốt dầm", "t2f",
                 f"=IF(AND(TS_nd>=10,{nrow}>1),1,IF({nf}-{k1}>=2,1,0))",
                 ">= 10 thanh (cắt so le): thanh góc dài suốt -> 1 ; còn lại: = 1 khi số thanh dài suốt còn lại "
                 "sau hàng ngoài >= 2 (đặt ở 2 đầu hàng 2)", "")
    t2y = line("", "Số vị trí đai phụ giữ lớp 2 - 1 dầm theo Y", "t2Y",
               f"=IF({has2}=0,0,IF({full2}=1,{tiy},2*{STA(f'MIN(TS_B/2-TS_cv,TS_B/2-{cuty})-(TS_L2/2+{wq})', 'TS_sp/1000')}"
               f"+{STA(f'2*MIN(TS_L2/2-{wq},TS_B/2-{cuty})', 'TS_sp/1000')}))",
               "Trong đoạn thanh ngắt (±(B/2 − cutY)), bỏ vùng giao dầm", "vị trí")
    t2x = line("", "Số vị trí đai phụ giữ lớp 2 - 1 dầm theo X", "t2X",
               f"=IF({has2}=0,0,IF({full2}=1,{tix},2*{STA(f'MIN(TS_L/2-TS_cv,TS_L/2-{cutx})-(TS_L1/2+{wq})', 'TS_sp/1000')}"
               f"+{STA(f'2*MIN(TS_L1/2-{wq},TS_L/2-{cutx})', 'TS_sp/1000')}))",
               "Trong đoạn thanh ngắt (±(L/2 − cutX)), bỏ vùng giao dầm", "vị trí")
    hh_ = line("", "Chiều cao toàn móng", "H", "=TS_H1+TS_H2", "H1 + H2", "m")
    n20 = line("", "Số đai cổ móng (1 cổ)", "n20", f"=1+MAX(1,ROUNDUP(({hh_}-TS_cv-TS_hd)/(TS_sh/1000)-0.000002,0))",
               "1 + ROUNDUP((H − cv − h)/@)", "đai")
    n21 = line("", "Số mức đai phụ cổ móng (1 cổ)", "nph", f"=1+MAX(1,ROUNDUP(({hh_}-TS_cv-TS_hd)/(TS_sph/1000)-0.000002,0))",
               "1 + ROUNDUP((H − cv − h)/@ đai phụ)", "mức")
    mcm = line("", "Số thanh thép chủ cổ móng 1 cạnh", "mcm", "=1+INT(TS_ncm/4)", "1 + n/4", "thanh")
    scm = line("", "Khoảng cách thanh chủ cổ móng", "scm", f"=1000*(TS_bc-2*TS_cv)/MAX(1,{mcm}-1)",
               "(b − 2cv)/(m − 1)", "mm")
    gst = line("", "Bước thanh được đai phụ giữ", "gst", f"=MAX(1,MIN(2,INT(400/{scm})))",
               "Khoảng cách thanh được giữ <= 400 mm: 1 (mọi thanh) hoặc 2 (cách 1 thanh)", "")
    none_ = f"OR({mcm}<=2,AND(1000*(TS_bc-2*TS_cv)<=400,{mcm}<=4))"
    nst21 = line("", "Số đai phụ thẳng mỗi mức (mỗi phương)", "n21",
                 f"=IF({none_},0,VLOOKUP({mcm},TS_DPC,IF({gst}=1,2,4),FALSE))",
                 "Tra bảng BỐ TRÍ ĐAI PHỤ CỔ MÓNG (sheet THONG SO) theo m và g", "thanh")
    gmax = line("", "Chỉ số thanh xa nhất của đai bát giác (−1: không có)", "gmax",
                f"=IF({none_},-1,VLOOKUP({mcm},TS_DPC,IF({gst}=1,3,5),FALSE))", "Tra bảng như trên", "")
    rr = "MAX(0.35,0.5*TS_dcm/1000*40)"
    hb = f"(40*(TS_bc/2-TS_cv)-0.2-{rr})"
    rt = f"({rr}+0.12)"
    oa = line("", "Đai bát giác: nửa cạnh thẳng", "ao",
              f"=IF({gmax}<0,0,(({gmax}-({mcm}-1)/2)*2*{hb}/({mcm}-1)+0.41421356*{rt})/40)",
              "Theo hình mặt cắt 5-5 (bo góc bán kính thép chủ)", "m")
    of = line("", "Đai bát giác: nửa bề rộng", "fo", f"=IF({gmax}<0,0,({hb}+{rt})/40)", "Theo hình mặt cắt 5-5", "m")
    grp = True
    nn23 = line("", "Số thanh lưới đầu trụ mỗi phương, mỗi lớp", "n23",
                "=1+ROUNDUP((TS_bc-2*TS_cv)/(TS_sl/1000)-0.000002,0)", "1 + ROUNDUP((b − 2cv)/@)", "thanh")
    ntl = line("", "Số bậc thang leo (1 trụ)", "ntl",
               "=IF(TS_H2<=TS_htl+0.000001,0,ROUND((TS_H2-TS_htl-TS_cv)/(TS_stl/1000),0)+1)",
               "H2 <= 1.5 m: 0 ; còn lại ROUND((H2 − 1.5 − cv)/@)+1", "bậc")

    # công thức từng vị trí: (Ø, chiều dài 1 thanh (m), số lượng, diễn giải chiều dài, diễn giải số lượng)
    lg = "(TS_hd-2*TS_cv)"
    l7 = "(SQRT(2)*(2*TS_a+TS_e)-2*SQRT(2)*TS_cv-2*TS_cv)"
    l8 = "(SQRT(2)*(TS_a+TS_e)-2*TS_cv)"
    n8 = f"MAX(1,INT(({l8}+2*TS_cv)/0.3-1+0.000000001))"
    n9 = "MAX(1,INT((TS_hd-TS_hs)/0.2+1+0.000000001))"
    F = {
        1: ("TS_ds", "TS_B-2*TS_cv", f"2*{nx}", "B − 2cv", "2 lớp × nX"),
        2: ("TS_ds", "TS_B/2-TS_c/2-2*TS_cv", f"4*{nh}", "B/2 − c/2 − 2cv", "2 lớp × 2 bên lỗ × nh"),
        3: ("TS_ds", "TS_L-2*TS_cv", f"2*{ny}", "L − 2cv", "2 lớp × nY"),
        4: ("TS_ds", "TS_L/2-TS_c/2-2*TS_cv", f"4*{nh}", "L/2 − c/2 − 2cv", "2 lớp × 2 bên lỗ × nh"),
        5: ("TS_ds", "ROUND(TS_c*SQRT(2),2)+0.8", f"IF(MOD(4*{nh},8)=0,4*{nh},INT(4*{nh}/8)*8)",
            "c·√2 + 800", "4·nh làm tròn xuống bội số 8 (2 lớp × 4 góc)"),
        6: ("TS_dch", "TS_hs-TS_cv-TS_cvb-2*TS_ds/1000+2*TS_lch", f"{chx}*{chy}-{chh}",
            "Thân (h1 − cv − cvb − 2Ø bản) + 2 chân 100", "chX × chY − thanh trong lỗ"),
        7: ("TS_dt", f"{l7}+2*{lg}", "IF(TS_e>0.000001,16,0)", "√2(2a + e) − 2√2·cv − 2cv + 2 chân (h − 2cv)", "4 thanh × 4 chỗ giao dầm"),
        8: ("TS_dt", f"{l8}+2*{lg}", f"IF(TS_e>0.000001,2*{n8}*4,0)", "√2(a + e) − 2cv + 2 chân (h − 2cv)",
            "2 họ × INT((l8 + 2cv)/300 − 1) × 4 chỗ giao"),
        9: ("TS_dt", f"4*{l8}+2*TS_hk", f"IF(TS_e>0.000001,{n9}*4,0)", "4·l8 + 2 móc", "INT((h − h1)/200 + 1) × 4 chỗ giao"),
        10: ("TS_dm", "TS_B-2*TS_cv", f"{nf}*4", "B − 2cv", "nf × 2 mặt × 2 dầm"),
        11: ("TS_dm", f"TS_B-2*{cuty}", f"{ncut}*4", "B − 2·cutY", "nc × 2 mặt × 2 dầm"),
        12: ("TS_ds", "TS_B-2*TS_cv", f"{ngia}*4", "B − 2cv", "ngia × 2 bên × 2 dầm"),
        13: ("TS_dm", "TS_L-2*TS_cv", f"{nf}*4", "L − 2cv", "nf × 2 mặt × 2 dầm"),
        14: ("TS_dm", f"TS_L-2*{cutx}", f"{ncut}*4", "L − 2·cutX", "nc × 2 mặt × 2 dầm"),
        15: ("TS_ds", "TS_L-2*TS_cv", f"{ngia}*4", "L − 2cv", "ngia × 2 bên × 2 dầm"),
        16: ("TS_dt", "2*(TS_a-2*TS_cv)+2*(TS_hd-2*TS_cv)+2*TS_hk", f"2*({sty}+{stx})",
             "2(a − 2cv) + 2(h − 2cv) + 2 móc", "2 dầm × (nstY + nstX)"),
        17: ("TS_dp", "TS_a-2*TS_cv+2*TS_hk", f"2*({tiy}+{tix})*{ngia}+4*({t2y}+{t2x})",
             "a − 2cv + 2 móc", "2 dầm × (ntY + ntX) × ngia + 2 dầm × 2 mặt × (t2Y + t2X)"),
        18: ("TS_dp", "TS_hd-2*TS_cv+2*TS_hk", f"IF(TS_a>0.600001,2*({tiy}+{tix}),0)", "h − 2cv + 2 móc", "2 dầm × (ntY + ntX)"),
        19: ("TS_dcm", f"{hh_}-TS_cv-TS_cvb-2*TS_ds/1000+TS_lneo", "TS_ncm*TS_n",
             "H − cv − cvb − 2Ø bản + neo 500", "số thanh × số trụ"),
        20: ("TS_dh", "4*(TS_bc-2*TS_cv)+2*TS_hk", f"{n20}*TS_n", "4(b − 2cv) + 2 móc", "n20 × số trụ"),
        21: ("TS_dph", "TS_bc-2*TS_cv+2*TS_hk", f"{n21}*2*{nst21}*TS_n", "b − 2cv + 2 móc",
             "nph × 2 phương × n21 × số trụ"),
        22: ("TS_dph", "IF({gm}<0,0,8*{a}+4*SQRT(2)*({f}-{a})+2*TS_hk)", f"IF({gmax}<0,0,{n21}*TS_n)", "8·ao + 4√2(fo − ao) + 2 móc",
             "nph × số trụ"),
        23: ("TS_dl", "TS_bc-2*TS_cv", f"TS_nl*2*{nn23}*TS_n", "b − 2cv", "số lớp × 2 phương × n23 × số trụ"),
        24: ("TS_dtl", "2*TS_ltl+TS_wtl", f"{ntl}*2", "2 chân 350 + bậc 300", "ntl × 2 trụ có thang"),
    }
    if grp:
        f22 = F[22]
        F[22] = (f22[0], f22[1].format(a=oa, f=of, gm=gmax), *f22[2:])

    sec("E. CỐT THÉP (chiều dài, số lượng tính bằng công thức - xem cột I, J)")
    r += 1
    for i, h in enumerate(["Vị trí", "Hình dạng / ghi chú", "Ø (mm)", "Chiều dài 1 thanh (mm)", "Số lượng",
                           "kg/m", "Tổng dài (m)", "KL (kg)", "Cách tính chiều dài", "Cách tính số lượng"], 1):
        _cell(wk, r, i, h, bold=True, fill=HEAD, align=Alignment(horizontal="center", vertical="center", wrap_text=True))
    desc = {r["pos"]: r["name"] for r in g.kl}
    desc.update({k: v for k, v in NAMES.items() if k not in desc})
    rs0 = r + 1
    for pos in sorted(F):
        r += 1
        fd, fl, fq, tl, tq = F[pos]
        # số hiệu chạy theo số liệu: chỉ đánh số các thanh có số lượng > 0 (giống bảng thống kê trên bản vẽ)
        _cell(wk, r, 1, f'=IF(E{r}>0,COUNTIF(E${rs0}:E{r},">0"),"—")', align=Alignment(horizontal="center"))
        _cell(wk, r, 2, desc.get(pos, ""))
        _cell(wk, r, 3, f"={fd}", "0")
        _cell(wk, r, 4, f"=ROUND(1000*({fl}),0)", "0")
        _cell(wk, r, 5, f"={fq}", "0")
        _cell(wk, r, 6, f"=IFERROR(VLOOKUP(C{r},TS_KGM,2,FALSE),C{r}^2/162)", "0.000")
        _cell(wk, r, 7, f"=D{r}/1000*E{r}", "0.00")
        _cell(wk, r, 8, f"=G{r}*F{r}", "0.00")
        _cell(wk, r, 9, tl)
        _cell(wk, r, 10, tq)
    rs1 = r
    r += 1
    _cell(wk, r, 2, "TỔNG CỘNG THÉP", bold=True)
    _cell(wk, r, 8, f"=SUM(H{rs0}:H{rs1})", "0.00", bold=True)
    rtot = r
    r += 2
    _cell(wk, r, 2, "Phân loại theo đường kính", bold=True, fill=HEAD)
    _cell(wk, r, 3, "Ø (mm)", bold=True, fill=HEAD)
    _cell(wk, r, 7, "Tổng dài (m)", bold=True, fill=HEAD)
    _cell(wk, r, 8, "KL (kg)", bold=True, fill=HEAD)
    rdia = {}
    for d, _, _ in g.klc:
        r += 1
        rdia[int(d + 0.5)] = r
        _cell(wk, r, 2, f"Thép Ø{int(d + 0.5)}")
        _cell(wk, r, 3, int(d + 0.5))
        _cell(wk, r, 7, f"=SUMIF($C${rs0}:$C${rs1},C{r},$G${rs0}:$G${rs1})", "0.00")
        _cell(wk, r, 8, f"=SUMIF($C${rs0}:$C${rs1},C{r},$H${rs0}:$H${rs1})", "0.00")
    r += 2
    _cell(wk, r, 2, "Bu lông neo", bold=True)
    _cell(wk, r, 4, f"{model.bolt_text(p)}  ({p['BLN']})")
    wk.freeze_panes = "A4"
    _summary(wb, p, g, rows, rdia, rtot)
    wb.active = 0
    wb.save(path)


def _summary(wb, p, g, rows, rdia, rtot):
    """Sheet TONG KE: bảng tổng kê khối lượng 1 vị trí móng, mỗi số liệu 1 ô (theo mẫu)."""
    ws = wb.create_sheet("TONG KE", 0)
    K = "'KHOI LUONG'!"
    cols = [(f"{p['BTM']} (m³)", f"={K}{rows['Vbt']}", "0.00"),
            (f"{p['BTL']} (m³)", f"={K}{rows['Vlot']}", "0.00"),
            ("Vữa M10 (m³)", f"={K}{rows['Vvua']}", "0.000")]
    for d in sorted(rdia):
        cols.append((f"Ø{d}", f"={K}$H${rdia[d]}", "0.00"))
    cols += [("TỔNG K.LƯỢNG THÉP (kg)", f"={K}$H${rtot}", "0.00")]
    if p["HC"] > 1e-6:
        cols += [("K.LƯỢNG ĐỆM CÁT (m³)", f"={K}{rows['Vcat']}", "0.00"),
                 ("DIỆN TÍCH VẢI ĐỊA (m²)", f"={K}{rows['Svai']}", "0.00")]
    cols += [("BULONG NEO", model.bolt_text(p), None),
             ("Đào đất (m³)", f"={K}{rows['Vdao']}", "0.00"),
             ("Đắp đất (m³)", f"={K}{rows['Vdap']}", "0.00"),
             ("Ván khuôn (m²)", f"={K}{rows['VK']}", "0.00"),
             ("Rộng đền bù (m)", "=TS_B+2*0.5", "0.00"),
             ("Dài đền bù (m)", "=TS_L+2*0.5", "0.00"),
             ("Diện đền bù (m²)", None, "0.00"),
             ("Chiều sâu hố đào (m)", f"={K}{rows['Hd']}", "0.00")]
    n = len(cols)
    ws.page_setup.orientation = "landscape"           # in vừa 1 trang ngang
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 64
    ws.row_dimensions[3].height = 24
    c = ws.cell(row=1, column=1, value="BẢNG TỔNG KÊ KHỐI LƯỢNG CHO 01 VỊ TRÍ MÓNG")
    c.font = Font(size=14)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    for j in range(1, n + 1):
        ws.cell(row=1, column=j).border = BOX
    from openpyxl.utils import get_column_letter as L
    for j, (h, f, fmt) in enumerate(cols, 1):
        ws.column_dimensions[L(j)].width = 11 if len(h) < 14 else 12.5
        _cell(ws, 2, j, h, align=Alignment(horizontal="center", vertical="center", wrap_text=True))
        if h.startswith("Diện đền bù"):
            f = f"={L(j - 2)}3*{L(j - 1)}3"          # rộng đền bù x dài đền bù
        _cell(ws, 3, j, f, fmt, align=Alignment(horizontal="center", vertical="center"))
