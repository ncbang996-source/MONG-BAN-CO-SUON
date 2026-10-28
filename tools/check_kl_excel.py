"""Đối chiếu Excel khối lượng (công thức) với bảng thống kê thép của chương trình.

Xuất _KL.xlsx cho nhiều bộ thông số, cho LibreOffice tính lại công thức (soffice --headless),
rồi so Ø, chiều dài, số lượng từng vị trí thép.   python tools/check_kl_excel.py [thư_mục_tạm]
"""
import os
import subprocess
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
import openpyxl  # noqa: E402

from mongban import kl_excel, model  # noqa: E402
from mongban.geom import Geo  # noqa: E402

OUT = sys.argv[1] if len(sys.argv) > 1 else tempfile.mkdtemp()
os.makedirs(OUT, exist_ok=True)
BIG = dict(L1=6, L2=6, B=13.5, L=15, A=0.6, HD=1.2, HS=0.5, C=1.5, E=0.5, BB=1.0, H1=2.5, H2=4.5, HT=1, RD_N=10,
           HC=0.7, RS_S2=140, RS_S1=200, RCH_S2=420, RCH_S1=600, RC_N=32, MPIT=0.6)
CASES = {"mau": {}, "lon": BIG, "n7": dict(RD_N=7), "n10": dict(RD_N=10), "n8day": dict(RD_ARR=1),
         "kolo": dict(C=0.0), "kovat": dict(E=0.0), "h07": dict(HD=0.7), "a07": dict(A=0.7, BB=0.8),
         "rc32": dict(RC_N=32, BB=1.0), "h2thap": dict(H2=1.4), "n11": dict(RD_N=11), "n4": dict(RD_N=4),
         "n5": dict(RD_N=5), "n9": dict(RD_N=9), "n12": dict(RD_N=12), "n13": dict(RD_N=13), "n14": dict(RD_N=14),
         "n16": dict(RD_N=16), "n18": dict(RD_N=18), "n12a7": dict(RD_N=12, A=0.7, BB=0.8),
         "hep10": dict(RD_N=10, A=0.3, BB=0.5), "n16day": dict(RD_N=16, RD_ARR=1), "n11day": dict(RD_N=11, RD_ARR=1),
         "dt1": dict(HT=1, HDT=0.5), "dt2": dict(HT=1, HDT=1.0, HDD=0.5), "dt3": dict(HT=2, HDT=0.8, DT_MODE=1, DT_SEL="136"),
         "dt4": dict(HT=2, HDT=0.6, B=13, L=10, L1=4, L2=5, DT_MODE=1, DT_SEL="2456"), "dt5": dict(HT=0, HDT=0.5),
         "dt6": dict(HT=1, HDT=0.9, H1=1.0, BB=1.2, E=0.1), "dd1": dict(HDD=0.8), "dtlon": dict(BIG, HDT=1.0, HDD=0.8)}
# khối lượng bê tông / đất: ký hiệu ở cột C của sheet KHOI LUONG -> hàm trong model
VOLS = {"Vbt": model.vol_mong, "Vlot": model.vol_lot, "Vcat": model.vol_cat, "Vdao": model.vol_dao,
        "Vdap": model.vol_dap, "Vdd": model.vol_dd, "Vthua": model.vol_thua, "Vdt": model.vol_dt,
        "Vdtb": lambda p: model.vol_dt(p, below=True), "VK": model.area_vk, "Svai": model.area_vai}
base, _ = model.read_excel(os.path.join(ROOT, "samples", "D122-42_MB4-10x12-2.5.xlsm"))
geos = {}
params = {}
for nm, ov in CASES.items():
    p = dict(base)
    p.update({k: (v if isinstance(v, str) else float(v)) for k, v in ov.items()})
    params[nm] = p
    p["SCALE"] = model.auto_scale(p)
    g = Geo(p)
    g.calc_kl()
    kl_excel.export(p, g, os.path.join(OUT, f"{nm}.xlsx"))
    geos[nm] = g
calc = os.path.join(OUT, "calc")
subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", calc] +
               [os.path.join(OUT, f"{nm}.xlsx") for nm in CASES], capture_output=True, timeout=900)
bad = 0
for nm, g in geos.items():
    ws = openpyxl.load_workbook(os.path.join(calc, f"{nm}.xlsx"), data_only=True)["KHOI LUONG"]
    got = {row[1].value: (row[2].value, row[3].value, row[4].value) for row in ws.iter_rows()
           if row[1].value in kl_excel.NAMES.values() and isinstance(row[4].value, (int, float))}
    names = {**kl_excel.NAMES, **{r["pos"]: r["name"] for r in g.kl}}
    exp = {r["pos"]: r for r in g.kl}
    errs = []
    for pos, nmp in names.items():
        d, ln, q = got.get(nmp, (None, None, None))
        e = exp.get(pos)
        if q is None:
            errs.append(f"{pos}: thiếu dòng / lỗi công thức")
        elif e and (abs(ln - e["len"]) > 1.0 or int(round(q)) != int(e["n"]) or int(d) != int(e["dia"] + 0.5)):
            errs.append(f"{pos}: Excel Ø{d} L={ln} n={q} | chương trình Ø{e['dia']} L={e['len']:.0f} n={e['n']}")
        elif not e and q:
            errs.append(f"{pos}: Excel n={q}, chương trình không có")
    vals = {row[2].value: row[4].value for row in ws.iter_rows() if row[2].value in VOLS}
    for sym, fn in VOLS.items():
        want = fn(params[nm])
        if sym not in vals or vals[sym] is None or abs(float(vals[sym]) - want) > 0.02:
            errs.append(f"{sym}: Excel {vals.get(sym)} | chương trình {want:.3f}")
    print(f"{nm:8s} {'OK' if not errs else 'SAI'}")
    for x in errs:
        print("    ", x)
    bad += len(errs)
print("Tổng số lỗi:", bad)
sys.exit(1 if bad else 0)
